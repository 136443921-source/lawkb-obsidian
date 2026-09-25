#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
LTI 交付门禁脚本  v2（2026-09-01 升级）

变更说明（v1 → v2）：
  v1 把 docx 全文拼成单个字符串后一次性校验，导致「上下文相关」类规则
  （典型如 R102 条号↔内容一致性）产生结构性误报：全文所有法条号与全文
  所有主题关键词被强行放进同一上下文，例如文中同时出现「上诉」与民诉法
  第一百五十条（法庭笔录），即被判为「上诉应对应第171条而非第150条」。
  实测同一份文书全文跑报 R102、逐段跑 0 命中，确认为拼接误报。

  v2 采用双粒度校验：
    · 段落粒度（主）：逐段 / 逐表格单元格校验，命中上下文类规则更准确；
    · 全文粒度（辅）：仅采信跨段一致性类规则（C301/C302/C304）与条号区间
      类规则（R101）等，用于捕捉跨段落的金额、名称、案号冲突；
      其余规则在全文中命中一律降级为提示，不再阻断。

用法: python3 lti_gate.py <docx路径...>
  PASS  : 打印提示汇总（WARNING/INFO 逐条列出），exit 0
  REJECT: 打印阻断原因，exit 1（阻断交付）
"""
import sys
import os
import re

LTI_REF = os.path.expanduser("~/.workbuddy/skills/LTI文本监控器/references")
sys.path.insert(0, LTI_REF)

from main import wrap_output, LTIRejectError  # noqa: E402
from yuandian_verify import load_case_cache, extract_case_requests, record_case_verification  # noqa: E402
from yuandian_rest_client import api_key_ready, verify_cases_rest  # noqa: E402  (Level3 全自动：元典 REST 直连)
import docx  # noqa: E402

# P2（2026-09-03）：案例库存在性核验缓存，接入 T502 门禁（堵「案例存在性」硬边界）
# 编造案号(exists=False)→REJECT；未核验→INFO（提示先 verify_cases.py 核验）；真实→静默
try:
    CASE_VERIFY = load_case_cache()
except Exception:
    CASE_VERIFY = {}

# 全文粒度下仍然可信的规则（跨段一致性 / 条号区间等，与上下文无关）
FULLTEXT_TRUSTED = {"C301", "C302", "C304", "R101", "R103", "R104",
                    "R105", "R106", "R108"}

RULE_ID_RE = re.compile(r"\[([A-Z]\d{3})\]")


def extract_blocks(path):
    """返回 [(来源标签, 文本)]，保留段落与表格单元格边界。"""
    d = docx.Document(path)
    blocks = []
    for i, p in enumerate(d.paragraphs):
        if p.text.strip():
            blocks.append((f"段落[{i}]", p.text.strip()))
    for ti, tb in enumerate(d.tables):
        for ri, row in enumerate(tb.rows):
            for ci, c in enumerate(row.cells):
                if c.text.strip():
                    blocks.append((f"表格{ti}·行{ri}·列{ci}", c.text.strip()))
    return blocks


def extract_full(path):
    return "\n".join(t for _, t in extract_blocks(path))


def _parse_hints(out):
    """从 wrap_output 返回值中提取提示行。"""
    hints = []
    in_hint = False
    for line in out.split("\n"):
        s = line.strip()
        if s.startswith("【") or s.startswith("[") or s.startswith("💡"):
            in_hint = True
        if in_hint and s:
            hints.append(s)
    return [h for h in hints
            if not h.startswith(("## 【文本监控器参与说明", "1. 本文件由人工智能",
                                 "- ☐", "4. 责任声明"))]


def _rule_ids(text):
    return set(RULE_ID_RE.findall(text or ""))


def _dedup(items):
    """同一问题在段落级与全文级各报一次时，只保留首次（段落级标签更精确）。"""
    seen = set()
    out = []
    for it in items:
        m = RULE_ID_RE.search(it)
        if m:
            # 以「规则号 + 原因主体」为指纹，忽略来源前缀与换行差异
            body = re.sub(r"\s+", "", it[m.end():])
            key = (m.group(1), body[:80])
        else:
            key = re.sub(r"\s+", "", it)[:100]
        if key in seen:
            continue
        seen.add(key)
        out.append(it)
    return out


# ---- T502 自动预检（2026-09-06 自动化触发）：文书进门禁即自动提取案号→比对缓存→写待核验清单 ----
PENDING_PATH = os.path.join(LTI_REF, "provision_index", "pending_verify.json")


def write_pending(path, reqs):
    """把未核验案号写入待核验清单（写前备份）。reqs: list of (req_dict, status)"""
    import json
    import shutil
    import datetime
    data = {
        "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source_doc": path,
        "pending": [{"ref": r["ref"], "raw": r["raw"], "status": st} for r, st in reqs],
    }
    if os.path.exists(PENDING_PATH):
        bak = PENDING_PATH + ".bak-" + datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
        try:
            shutil.copy2(PENDING_PATH, bak)
        except Exception:
            pass
    with open(PENDING_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return len(reqs)


def t502_preflight(path, cache):
    """T502 自动预检：提取案号→比对缓存→分离 已核验(静默)/待核验。返回 (verified, pending)。

    设计：脚本层只做「提取+比对+写待核验」，联网核验回填由会话协办闭环完成
    （或 Level3：若 credentials.json 就绪则脚本直接调元典 REST 全自动，详见运维手册）。
    """
    full = extract_full(path)
    reqs = extract_case_requests(full)
    verified, pending = [], []
    for r in reqs:
        rec = cache.get(r["ref"])
        if rec and rec.get("exists") is True:
            verified.append(r["ref"])
        elif rec and rec.get("exists") is False:
            pending.append((r, "EXISTS_FALSE(编造候选)"))
        else:
            pending.append((r, "UNVERIFIED"))
    if pending:
        write_pending(path, pending)
    return verified, pending


def run(path):
    """返回 (status, detail, hints)"""
    blocks = extract_blocks(path)
    hits = []          # 段落粒度命中的提示
    rejects = []       # 阻断项

    # ---------- 1. 段落粒度（主校验） ----------
    for label, text in blocks:
        try:
            out = wrap_output(
                text,
                ai_participation=["法条检索与效力校验", "文书草拟"],
                case_verify=CASE_VERIFY,
            )
            for h in _parse_hints(out):
                hits.append(f"{label} {h}")
        except LTIRejectError as e:
            rejects.append(f"{label} {str(e).strip()}")

    # ---------- 2. 全文粒度（仅采信跨段类规则） ----------
    full = extract_full(path)
    try:
        out = wrap_output(
            full,
            ai_participation=["法条检索与效力校验", "文书草拟", "法理逻辑审查", "一致性校验"],
            case_verify=CASE_VERIFY,
        )
        for h in _parse_hints(out):
            ids = _rule_ids(h)
            if ids & FULLTEXT_TRUSTED:
                hits.append(f"全文 {h}")
    except LTIRejectError as e:
        detail = str(e)
        ids = _rule_ids(detail)
        if ids & FULLTEXT_TRUSTED:
            rejects.append(f"全文 {detail.strip()}")
        else:
            # 非跨段类规则：疑为拼接误报，降级为提示，不阻断
            hits.append(f"全文（降级提示·非阻断）{detail.strip()}")

    rejects = _dedup(rejects)
    hits = _dedup(hits)
    if rejects:
        return "REJECT", "\n\n".join(rejects), hits
    return "PASS", "", hits


def main():
    global CASE_VERIFY
    files = sys.argv[1:]
    if not files:
        print("用法: python3 lti_gate.py <docx路径...>")
        sys.exit(2)
    rejected = []
    print("=" * 72)
    for f in files:
        name = os.path.basename(f)
        # ---- T502 自动预检（提取案号→比对缓存→写待核验清单）----
        v, p = t502_preflight(f, CASE_VERIFY)
        # ---- Level3 全自动：元典 REST 直连回填（仅当有待核验且 YD_API_KEY 就绪）----
        if p and api_key_ready():
            print(f"\n🔧 Level3 全自动：调元典 REST 核验 {len(p)} 案号…")
            try:
                recs = verify_cases_rest([r["ref"] for r, _ in p])
                if recs:
                    n = record_case_verification(recs, backup=True)
                    CASE_VERIFY = load_case_cache()      # 回填后重载缓存
                    v, p = t502_preflight(f, CASE_VERIFY)  # 真实项转静默
                    print(f"   ✅ 回填 {n} 条真实案号；重跑静默 {len(v)} / 仍需复核 {len(p)}")
                else:
                    print(f"   ⚠️ REST 未命中真实案号（可能未收录），降级 Level2 会话协办")
            except Exception as e:
                print(f"   ⚠️ Level3 自动核验异常（{e}），降级 Level2 会话协办")
        # ---- 预检结果播报 ----
        if v or p:
            print(f"\n🔎 T502 自动预检 {name}：提取 {len(v)+len(p)} 案号 → 已核验静默 {len(v)} / 待核验 {len(p)}")
            for r, st in p:
                print(f"   · 待核验 [{st}] {r['ref']}  (原文 {r['raw']})")
            if p:
                print(f"   ⚠️ 已写入待核验清单 → provision_index/pending_verify.json")
                if api_key_ready():
                    print(f"   🔧 Level3 已尝试直连；剩余项交会话协办调元典 rh_ptal_search 复核后重跑")
                else:
                    print(f"   🔧 会话协办：小强将自动调元典 rh_ptal_search 核验并回填 case_verify_cache.json，随后重跑门禁")
        status, detail, hints = run(f)
        if status == "REJECT":
            rejected.append(name)
            print(f"\n🔴 REJECT  {name}\n{detail}")
        else:
            print(f"\n✅ PASS  {name}（提示 {len(hints)} 条）")
            for h in hints[:12]:
                print(f"   · {h[:120]}")
    print("\n" + "=" * 72)
    if rejected:
        print(f"❌ 门禁阻断: {len(rejected)} 份未过 LTI → {rejected}")
        print("   必须修正后重跑通过方可交付。")
        sys.exit(1)
    print("✅ LTI 门禁全部通过（无 REJECT），可交付。")
    sys.exit(0)


if __name__ == "__main__":
    main()
